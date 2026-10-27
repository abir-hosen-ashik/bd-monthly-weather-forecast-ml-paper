import json
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from common import DATA, EXISTING, FIGURES, RESULTS, STATION_DIVISION

CACHE = DATA / "_eda_cache"
OUT = RESULTS / "historical"
OUT.mkdir(parents=True, exist_ok=True)
DIVS = ["Barisal", "Chittagong", "Dhaka", "Khulna", "Mymensingh", "Rajshahi", "Rangpur", "Sylhet"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
SEASONS6 = ["Summer", "Monsoon", "Autumn", "Late Autumn", "Winter", "Spring"]
ALIASES = {"Chittagong(Air-port)": "Chittagong", "Chittagong(City - Ambagan)": "Ambagan(ctg)", "M. Court": "M.Court",
           "Maijdee Court": "M.Court", "Sitakunda": "Sitakund", "Srimongal": "Srimangal", "Syedpur": "Sydpur",
           "Cox's": "Cox's Bazar", "Ctg Ambagan": "Ambagan(ctg)"}
plt.rcParams.update({"font.size": 10, "font.weight": "bold", "axes.labelweight": "bold", "axes.titleweight": "bold",
                     "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": 200, "savefig.bbox": "tight"})


def season6(month, day):
    md = month * 100 + day
    if 414 <= md <= 613:
        return "Summer"
    if 614 <= md <= 814:
        return "Monsoon"
    if 815 <= md <= 1013:
        return "Autumn"
    if 1014 <= md <= 1213:
        return "Late Autumn"
    if md >= 1214 or md <= 213:
        return "Winter"
    return "Spring"


def rain_season(m):
    return "Summer" if m in (3, 4, 5) else ("Rainy Season" if m in (6, 7, 8, 9, 10) else "Winter")


def division_map(version):
    old = json.loads((CACHE / "divisions_temp99.json").read_text())
    mp = {s: d for d, sts in old.items() for s in sts}
    mp.update({"Dinajpur": "Rangpur", "Ambagan(ctg)": "Chittagong"})
    if version == "qc":
        mp["Ishurdi"] = "Rajshahi"
    for s, d in STATION_DIVISION.items():
        mp.setdefault(s, d)
    return mp


SOURCES = {
    "t99": ("Temp,Rainfall&Humidity(1999-2022)-Latest/Temp,Rainfall&Humidity(1999-2022)/Temperature/All Station(1999-2022).xlsx",
            ["sn", "Station", "Year", "Month", "Day", "Tmax", "Tmin", "Tavg"]),
    "r99": ("Temp,Rainfall&Humidity(1999-2022)-Latest/Temp,Rainfall&Humidity(1999-2022)/Rainfall/Rainfall All Station(1999-2022).xlsx",
            ["sn", "Station", "Year", "Month", "Rainfall"]),
    "h99": ("Temp,Rainfall&Humidity(1999-2022)-Latest/Temp,Rainfall&Humidity(1999-2022)/Humidity/Humidity All Station(1999-2022).xlsx",
            ["sn", "Station", "Year", "Month", "Day", "Humidity"]),
    "t80": ("Temp,Rain and Humidity Data(1980-2022)-Latest/Temp,Rain and Humidity Data(1980-2022)/Temperature 1980-2022.xlsx",
            ["sn", "Station", "Year", "Month", "Day", "Tmax", "Tmin", "Tavg"]),
    "r80": ("Temp,Rain and Humidity Data(1980-2022)-Latest/Temp,Rain and Humidity Data(1980-2022)/Rainfall 1980-2022.xlsx",
            ["sn", "Station", "Year", "Month", "Rainfall"]),
    "h80": ("Temp,Rain and Humidity Data(1980-2022)-Latest/Temp,Rain and Humidity Data(1980-2022)/Humidity 1980-2022.xlsx",
            ["sn", "Station", "Year", "Month", "Day", "Humidity"]),
}
DIV_FILES = "Temp,Rainfall&Humidity(1999-2022)-Latest/Temp,Rainfall&Humidity(1999-2022)/Temperature"


def build_cache():
    import openpyxl
    CACHE.mkdir(parents=True, exist_ok=True)
    for k, (rel, cols) in SOURCES.items():
        if (CACHE / f"{k}.csv.gz").exists():
            continue
        wb = openpyxl.load_workbook(EXISTING / rel, read_only=True)
        rows = [r[:len(cols)] for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if r[1] is not None]
        x = pd.DataFrame(rows, columns=cols).drop(columns="sn")
        for c in x.columns:
            if c != "Station":
                x[c] = pd.to_numeric(x[c], errors="coerce")
        x["Station"] = x["Station"].astype(str).str.strip()
        x.to_csv(CACHE / f"{k}.csv.gz", index=False)
    if not (CACHE / "divisions_temp99.json").exists():
        mp = {}
        for f in sorted((EXISTING / DIV_FILES).glob("* Division.xlsx")):
            wb = openpyxl.load_workbook(f, read_only=True)
            mp[f.name.replace(" Division.xlsx", "")] = sorted({str(r[1]).strip() for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if r[1]})
        (CACHE / "divisions_temp99.json").write_text(json.dumps(mp))


def load(version):
    build_cache()
    d = {}
    for k in ["t99", "t80", "h99", "h80", "r99", "r80"]:
        x = pd.read_csv(CACHE / f"{k}.csv.gz")
        x["Station"] = x["Station"].str.strip().replace(ALIASES)
        d[k] = x
    qc = {}
    for k in ["t99", "t80"]:
        t = d[k]
        if version == "qc":
            bad = (t.Tmax < 0) | (t.Tmax > 48) | (t.Tmin < -2) | (t.Tmin > 40) | (t.Tmin > t.Tmax) | (t.Tmax == 0) | (t.Tmin == 0)
            t.loc[bad, ["Tmax", "Tmin"]] = np.nan
            qc[k] = int(bad.sum())
        t["Tavg"] = (t.Tmax + t.Tmin) / 2
    for k in ["h99", "h80"]:
        h = d[k]
        if version == "qc":
            bad = (h.Humidity < 5) | (h.Humidity > 100)
            h.loc[bad, "Humidity"] = np.nan
            qc[k] = int(bad.sum())
    mp = division_map(version)
    for k, x in d.items():
        x["Division"] = x["Station"].map(mp)
        if "Day" in x:
            x["Season"] = [season6(m, dd) for m, dd in zip(x.Month, x.Day)]
        else:
            x["RSeason"] = x.Month.map(rain_season)
            x["Season"] = x.Month.map({1: "Winter", 2: "Spring", 3: "Spring", 4: "Summer", 5: "Summer", 6: "Monsoon",
                                       7: "Monsoon", 8: "Autumn", 9: "Autumn", 10: "Late Autumn", 11: "Late Autumn",
                                       12: "Winter"})
    return d, qc


def annotate_line(ax, x, y, idx, color, fmt="{:.2f}"):
    for i in idx:
        ax.annotate(fmt.format(y.iloc[i]), (x.iloc[i], y.iloc[i]), textcoords="offset points", xytext=(4, 6),
                    fontsize=8, color="black")


def yearly_tmm(ax, t, title=None):
    y = t.groupby("Year")[["Tmin", "Tmax", "Tavg"]].mean()
    for c, col, lab in [("Tmin", "blue", "Minimum Temperature"), ("Tmax", "red", "Maximum Temperature"),
                        ("Tavg", "green", "Mean Temperature")]:
        s = y[c].reset_index(drop=True)
        ax.plot(y.index, y[c], color=col, lw=3, label=lab)
        mid = len(s) // 2
        annotate_line(ax, pd.Series(y.index), s, [0, mid, len(s) - 1], col)
    ax.set_xticks(y.index[::1 if len(y) < 30 else 2])
    ax.tick_params(axis="x", rotation=90, labelsize=7)
    ax.set_xlabel("Year")
    ax.set_ylabel("Temperature (°C)")
    ax.legend(fontsize=6, loc="best")
    if title:
        ax.set_title(title, fontsize=9)
    return y


def bars(ax, s, ylabel, xlabel, ylim=None, colors=None, rot=45):
    colors = colors or plt.get_cmap("tab10").colors
    b = ax.bar(range(len(s)), s.values, color=[colors[i % len(colors)] for i in range(len(s))], edgecolor="black")
    for rect, v in zip(b, s.values):
        ax.text(rect.get_x() + rect.get_width() / 2, v, f"{v:.2f}", ha="center", va="bottom", fontsize=7)
    ax.set_xticks(range(len(s)))
    ax.set_xticklabels(s.index, rotation=rot, ha="right" if rot else "center")
    ax.set_ylabel(ylabel)
    ax.set_xlabel(xlabel)
    if ylim:
        ax.set_ylim(*ylim)


def line_extremes(ax, s, ylabel, xlabel, unit="mm"):
    ax.plot(s.index, s.values, color="blue", lw=3)
    i, j = s.values.argmax(), s.values.argmin()
    ax.annotate(f"Max: {s.iloc[i]:.2f} {unit}", (s.index[i], s.iloc[i]), xytext=(10, 0), textcoords="offset points",
                color="blue", fontsize=8, arrowprops=dict(arrowstyle="->", color="blue"))
    ax.annotate(f"Min: {s.iloc[j]:.2f} {unit}", (s.index[j], s.iloc[j]), xytext=(10, 0), textcoords="offset points",
                color="red", fontsize=8, arrowprops=dict(arrowstyle="->", color="red"))
    ax.set_ylabel(ylabel)
    ax.set_xlabel(xlabel)


def save(fig, version, n, sub=None):
    d = FIGURES / "historical" / version
    d.mkdir(parents=True, exist_ok=True)
    fig.savefig(d / (f"fig{n:02d}.png" if sub is None else f"fig{n:02d}{sub}.png"))
    plt.close(fig)


def grid8(func, version, n, title):
    fig, axes = plt.subplots(4, 2, figsize=(10, 13))
    for ax, dv in zip(axes.flat, DIVS):
        func(ax, dv)
        ax.set_title(dv, fontsize=9)
    fig.suptitle(title, fontsize=11)
    fig.tight_layout()
    save(fig, version, n)
    for i, dv in enumerate(DIVS):
        fig, ax = plt.subplots(figsize=(6.5, 4.2))
        func(ax, dv)
        save(fig, version, n, f"_{i}")


def make_all(version):
    d, qc = load(version)
    t99, t80, h99, h80, r99, r80 = (d[k] for k in ["t99", "t80", "h99", "h80", "r99", "r80"])
    key = {}

    def counts(tt, rr, hh, n):
        fig, axes = plt.subplots(2, 2, figsize=(10, 8))
        sm = lambda x: x.groupby(["Station", "Year", "Month"]).ngroups if "Day" in x else len(x)
        tot = pd.Series({"Temperature": sm(tt), "Rainfall": sm(rr), "Humidity": sm(hh)})
        axes[0, 0].bar(tot.index, tot.values, color=["#2c3e50", "#c70039", "#3498db"])
        axes[0, 0].set_ylim(tot.min() * 0.995 if tot.max() / tot.min() < 1.05 else tot.min() * 0.8, tot.max() * 1.002)
        axes[0, 0].set_ylabel("Number of Data Point")
        for ax, x, lab in [(axes[0, 1], tt, "(b) Temperature"), (axes[1, 0], rr, "(c) Rainfall"), (axes[1, 1], hh, "(d) Humidity")]:
            c = (x.groupby(["Division", "Station", "Year", "Month"]).size().groupby("Division").size()
                 if "Day" in x else x.groupby("Division").size()).reindex(DIVS)
            ax.bar(c.index, c.values, color=plt.get_cmap("tab10").colors)
            ax.set_title(lab, fontsize=9)
            ax.tick_params(axis="x", rotation=45)
            ax.set_ylabel("Number of Data Point")
        axes[0, 0].set_title("(a) Total", fontsize=9)
        fig.tight_layout()
        save(fig, version, n)
        for i, (x, ttl) in enumerate([(None, None), (tt, "T"), (rr, "R"), (hh, "H")]):
            fig, ax = plt.subplots(figsize=(6.5, 4.7))
            if x is None:
                ax.bar(tot.index, tot.values, color=["#2c3e50", "#c70039", "#3498db"])
                ax.set_ylim(tot.min() * 0.995 if tot.max() / tot.min() < 1.05 else tot.min() * 0.8, tot.max() * 1.002)
                ax.set_xlabel("Weather Parameter")
            else:
                c = (x.groupby(["Division", "Station", "Year", "Month"]).size().groupby("Division").size()
                     if "Day" in x else x.groupby("Division").size()).reindex(DIVS)
                ax.bar(c.index, c.values, color=plt.get_cmap("tab10").colors)
                ax.set_xlabel("Division")
            ax.set_ylabel("Number of Data Point")
            ax.tick_params(axis="x", rotation=40)
            save(fig, version, n, f"_{i}")
        return tot
    key["count_99"] = counts(t99, r99, h99, 2).to_dict()
    key["count_80"] = counts(t80, r80, h80, 3).to_dict()
    for per, (tt, rr) in [("99", (t99, r99)), ("80", (t80, r80))]:
        key[f"divcount_T{per}"] = tt.groupby(["Division", "Station", "Year", "Month"]).size().groupby("Division").size().reindex(DIVS).to_dict()
        key[f"divcount_R{per}"] = rr.groupby("Division").size().reindex(DIVS).to_dict()
        key[f"nstations_T{per}"] = tt.Station.nunique()
        key[f"nstations_R{per}"] = rr.Station.nunique()
    oth = DATA / "_cache_other_daily.csv.gz"
    if oth.exists():
        o = pd.read_csv(oth)
        n = o.groupby(["Station", "Year", "Month"]).ngroups
        fig, ax = plt.subplots(figsize=(6, 4))
        ax.bar(["Pressure", "Windspeed", "Wind Direction", "Dew Point"], [n] * 4, color=["#2c3e50", "#9b59b6", "#3498db", "#c70039"])
        ax.set_ylim(n * 0.7, n * 1.05)
        ax.set_ylabel("Number of Data Point")
        ax.set_xlabel("Weather Parameter")
        ax.tick_params(axis="x", rotation=30)
        save(fig, version, 4)
        key["count_other"] = n

    fig, ax = plt.subplots(figsize=(7, 4.5))
    y = yearly_tmm(ax, t99)
    save(fig, version, 7)
    key["T99_yearly"] = y.round(2)
    grid8(lambda ax, dv: yearly_tmm(ax, t99[t99.Division == dv]), version, 8, "Yearly average temperature of eight divisions")
    for n, c, lab in [(9, "Tmax", "Maximum"), (10, "Tmin", "Minimum"), (11, "Tavg", "Mean")]:
        fig, ax = plt.subplots(figsize=(8, 5))
        g = t99.groupby(["Division", "Year"])[c].mean().unstack(0).reindex(columns=DIVS)
        for dv in DIVS:
            ax.plot(g.index, g[dv], label=dv, lw=1.5)
        overall = t99[c].mean()
        ax.axhline(overall, color="black", ls="--", label="Average Line")
        ax.text(g.index.max() + 0.5, overall, f"{overall:.2f} (°C)", fontsize=7, va="center")
        ax.set_title(f"Yearly Average {lab} Temperature of 8 Division")
        ax.set_xticks(g.index)
        ax.tick_params(axis="x", rotation=90, labelsize=7)
        ax.set_xlabel("Year")
        ax.set_ylabel("Temperature (°C)")
        ax.legend(fontsize=6)
        save(fig, version, n)
        key[f"T99_div_{c}"] = g.round(2)
    fig, ax = plt.subplots(figsize=(7, 4.5))
    m = t99.groupby("Month").Tavg.mean()
    m.index = MONTHS
    bars(ax, m, "Temperature", "Month", (10, 30.5))
    save(fig, version, 12)
    key["T99_monthly"] = m.round(2)
    grid8(lambda ax, dv: bars(ax, t99[t99.Division == dv].groupby("Month").Tavg.mean().set_axis([x[:3] for x in MONTHS]),
                              "Temperature", "Month", (10, 31), rot=0), version, 13, "Monthly average temperature of eight divisions")
    fig, ax = plt.subplots(figsize=(6, 4.5))
    s = t99.groupby("Season").Tavg.mean().reindex(SEASONS6)
    bars(ax, s, "Average Temperature", "Season", (15, 30))
    save(fig, version, 14)
    key["T99_seasonal"] = s.round(2)

    fig, ax = plt.subplots(figsize=(7, 4.5))
    key["T80_yearly"] = yearly_tmm(ax, t80).round(2)
    save(fig, version, 15)
    fig, ax = plt.subplots(figsize=(7, 4.5))
    m = t80.groupby("Month").Tavg.mean()
    m.index = MONTHS
    bars(ax, m, "Temperature", "Month", (10, 30))
    save(fig, version, 16)
    key["T80_monthly"] = m.round(2)
    fig, ax = plt.subplots(figsize=(6, 4.5))
    s = t80.groupby("Season").Tavg.mean().reindex(SEASONS6)
    bars(ax, s, "Average Temperature", "Season", (15, 30))
    save(fig, version, 17)
    key["T80_seasonal"] = s.round(2)

    for per, r, (fy, fdiv, fdall, fm, fdm, fs) in [("99", r99, (18, 19, 20, 21, 22, 23)), ("80", r80, (24, None, None, 25, None, 26))]:
        fig, ax = plt.subplots(figsize=(7, 4.5))
        yr = r.groupby("Year").Rainfall.mean()
        line_extremes(ax, yr, "Yearly Average Rainfall (mm)", "Year")
        ax.set_xticks(yr.index[::1 if per == "99" else 2])
        ax.tick_params(axis="x", rotation=90, labelsize=7)
        save(fig, version, fy)
        key[f"R{per}_yearly"] = yr.round(2)
        if fdiv:
            def ydiv(ax, dv):
                s = r[r.Division == dv].groupby("Year").Rainfall.mean()
                line_extremes(ax, s, "Rainfall (mm)", "Year")
                ax.tick_params(axis="x", labelsize=7)
            grid8(ydiv, version, fdiv, "Yearly average rainfall of eight divisions")
            fig, ax = plt.subplots(figsize=(8, 5))
            ax.plot(yr.index, yr.values, marker="o", ms=3, label="Bangladesh", lw=1.5)
            for dv in DIVS:
                s = r[r.Division == dv].groupby("Year").Rainfall.mean()
                ax.plot(s.index, s.values, marker="o", ms=3, label=dv, lw=1.2)
            ax.set_xticks(yr.index)
            ax.tick_params(axis="x", rotation=90, labelsize=7)
            ax.set_ylabel("Yearly Average Rainfall (mm)")
            ax.set_xlabel("Year")
            ax.legend(fontsize=6)
            save(fig, version, fdall)
            key[f"R{per}_div_yearly"] = r.groupby(["Division", "Year"]).Rainfall.mean().unstack(0).round(2)
        fig, ax = plt.subplots(figsize=(7, 4.5))
        mm = r.groupby("Month").Rainfall.mean()
        ax.plot(mm.index, mm.values, color="blue", lw=4)
        i = mm.values.argmax()
        ax.annotate(f"Max: {mm.iloc[i]:.2f} mm", (mm.index[i], mm.iloc[i]), xytext=(15, 10), textcoords="offset points",
                    color="red", arrowprops=dict(arrowstyle="->", color="red"))
        ax.set_xticks(range(1, 13))
        ax.set_xlabel("Month")
        ax.set_ylabel("Monthly Average Rainfall (mm)")
        save(fig, version, fm)
        key[f"R{per}_monthly"] = mm.round(2)
        if fdm:
            def mdiv(ax, dv):
                s = r[r.Division == dv].groupby("Month").Rainfall.mean()
                ax.plot(s.index, s.values, lw=3)
                i = s.values.argmax()
                ax.annotate(f"Max: {s.iloc[i]:.2f} mm", (s.index[i], s.iloc[i]), xytext=(10, 5), textcoords="offset points",
                            fontsize=7, color="blue")
                ax.set_xticks(range(1, 13))
                ax.set_ylabel("Rainfall (mm)")
            grid8(mdiv, version, fdm, "Monthly average rainfall of eight divisions")
        fig, ax = plt.subplots(figsize=(7, 4.5))
        sy = r.groupby(["RSeason", "Year"]).Rainfall.mean().unstack(0)
        for c, col in [("Rainy Season", "tab:orange"), ("Summer", "tab:blue"), ("Winter", "tab:green")]:
            ax.plot(sy.index, sy[c], lw=3, color=col)
            ax.text(sy.index.max() - 4, sy[c].iloc[-6], c, color="indigo", fontsize=8)
        ax.set_xlabel("Year")
        ax.set_ylabel("Seasonal Average Rainfall (mm)")
        save(fig, version, fs)
        key[f"R{per}_seasonal_max"] = sy.max().round(2)

    for per, h, (fy, fm, fs) in [("99", h99, (27, 28, 29)), ("80", h80, (30, 31, 32))]:
        fig, ax = plt.subplots(figsize=(7, 4.5))
        yh = h.groupby("Year").Humidity.mean()
        ax.plot(yh.index, yh.values, color="blue", lw=3)
        annotate_line(ax, pd.Series(yh.index), yh.reset_index(drop=True), [0, yh.values.argmax(), len(yh) - 1], "red")
        ax.set_xticks(yh.index)
        ax.tick_params(axis="x", rotation=90, labelsize=7)
        ax.set_xlabel("Year")
        ax.set_ylabel("Relative Humidity")
        save(fig, version, fy)
        key[f"H{per}_yearly"] = yh.round(2)
        fig, ax = plt.subplots(figsize=(7, 4.5))
        m = h.groupby("Month").Humidity.mean()
        m.index = MONTHS
        bars(ax, m, "Relative Humidity", "Month", (50, 90))
        save(fig, version, fm)
        key[f"H{per}_monthly"] = m.round(2)
        fig, ax = plt.subplots(figsize=(6, 4.5))
        s = h.groupby("Season").Humidity.mean().reindex(SEASONS6)
        bars(ax, s, "Average Humidity", "Season", (55, 90))
        save(fig, version, fs)
        key[f"H{per}_seasonal"] = s.round(2)

    def twin_yearly(n, a, alab, acol, b, blab, bcol, title):
        fig, ax = plt.subplots(figsize=(7, 4.5))
        ax.plot(a.index, a.values, color=acol, lw=2)
        ax.set_ylabel(alab, color=acol)
        ax2 = ax.twinx()
        ax2.plot(b.index, b.values, color=bcol, lw=2)
        ax2.set_ylabel(blab, color=bcol)
        ax.set_xticks(a.index)
        ax.tick_params(axis="x", rotation=45, labelsize=7)
        ax.set_xlabel("Year")
        ax.set_title(title)
        save(fig, version, n)

    def twin_bars(n, idx, tt, b, blab, title, xlab, lines=True):
        fig, ax = plt.subplots(figsize=(7, 4.5))
        x = np.arange(len(idx))
        if lines:
            for c, ls, lab in [("Tmax", "-", "Max Temperature"), ("Tmin", "--", "Min Temperature"), ("Tavg", "-.", "Avg Temperature")]:
                ax.plot(x, tt[c].values, ls, color="tab:red", marker=".", label=lab)
            ax.set_ylabel("Temperature (°C)", color="tab:red")
        else:
            ax.plot(x, tt.values, "-", color="tab:red", marker=".", label="Humidity")
            ax.set_ylabel("Relative Humidity (%)", color="tab:red")
        ax.legend(fontsize=6, loc="upper left")
        ax2 = ax.twinx()
        ax2.bar(x, b.values, color=plt.get_cmap("tab10").colors, alpha=0.7, label=blab)
        ax2.set_ylabel(blab, color="tab:blue")
        ax2.legend(fontsize=6, loc="upper right")
        ax.set_zorder(ax2.get_zorder() + 1)
        ax.patch.set_visible(False)
        ax.set_xticks(x)
        ax.set_xticklabels(idx, rotation=45)
        ax.set_xlabel(xlab)
        ax.set_title(title)
        save(fig, version, n)

    ty = t99.groupby("Year")[["Tmax", "Tmin", "Tavg"]].mean()
    tm = t99.groupby("Month")[["Tmax", "Tmin", "Tavg"]].mean()
    ts = t99.groupby("Season")[["Tmax", "Tmin", "Tavg"]].mean().sort_index()
    ry, rm, rs = r99.groupby("Year").Rainfall.mean(), r99.groupby("Month").Rainfall.mean(), r99.groupby("Season").Rainfall.mean().sort_index()
    hy, hm, hs = h99.groupby("Year").Humidity.mean(), h99.groupby("Month").Humidity.mean(), h99.groupby("Season").Humidity.mean().sort_index()
    mon3 = [x[:3] for x in MONTHS]
    twin_yearly(33, ty.Tavg, "Average Temperature (°C)", "tab:red", ry, "Rainfall (mm)", "tab:blue", "Yearly Average Mean Temperature and Rainfall")
    twin_bars(34, mon3, tm, rm, "Rainfall (mm)", "Monthly Temperature and Rainfall", "Month")
    twin_bars(35, list(ts.index), ts, rs.reindex(ts.index), "Rainfall (mm)", "Seasonal Temperature and Rainfall", "Season")
    twin_yearly(36, ty.Tavg, "Average Temperature (°C)", "tab:red", hy, "Relative Humidity (%)", "tab:blue", "Yearly Average Mean Temperature and Humidity")
    twin_bars(37, mon3, tm, hm, "Relative Humidity (%)", "Monthly Temperature and Humidity", "Month")
    twin_bars(38, list(ts.index), ts, hs.reindex(ts.index), "Relative Humidity (%)", "Seasonal Temperature and Humidity", "Season")
    twin_yearly(39, hy, "Relative Humidity (%)", "tab:red", ry, "Rainfall (mm)", "tab:blue", "Yearly Average Humidity and Rainfall")
    twin_bars(40, mon3, hm, rm, "Rainfall (mm)", "Monthly Humidity and Rainfall", "Month", lines=False)
    twin_bars(41, list(hs.index), hs, rs.reindex(hs.index), "Rainfall (mm)", "Seasonal Humidity and Rainfall", "Season", lines=False)
    key["T99_div_monthly"] = t99.groupby(["Division", "Month"]).Tavg.mean().unstack(0).round(2)
    key["T99_div_yearly_all"] = t99.groupby(["Division", "Year"])[["Tmax", "Tmin", "Tavg"]].mean().round(2)
    key["R99_div_monthly"] = r99.groupby(["Division", "Month"]).Rainfall.mean().unstack(0).round(2)
    key["R99_seasonal"] = r99.groupby(["RSeason", "Year"]).Rainfall.mean().unstack(0).round(2)
    key["R80_seasonal"] = r80.groupby(["RSeason", "Year"]).Rainfall.mean().unstack(0).round(2)
    key["T99_season_all"] = t99.groupby("Season")[["Tmax", "Tmin", "Tavg"]].mean().round(2)
    key["R99_season6"] = r99.groupby("Season").Rainfall.mean().round(2)
    key["H99_season_all"] = h99.groupby("Season").Humidity.mean().round(2)
    return key, qc


if __name__ == "__main__":
    versions = sys.argv[1:] or ["original", "qc"]
    allkeys = {}
    for v in versions:
        k, qc = make_all(v)
        allkeys[v] = k
        print(v, "done; QC-flagged days:", qc)
    pd.to_pickle(allkeys, OUT / "key_values.pkl")
